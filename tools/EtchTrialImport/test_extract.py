import unittest
from extract import parse_line,LAYOUT,parse_sheet
from openpyxl.utils import get_column_letter

class ExtractTests(unittest.TestCase):
    def fixture(self,line):
        start,ops,sc=LAYOUT[line];v={};f={}
        for br,ar,er in [(4,14,33),(9,19,38)]:
            for r in range(5):
                for c in range(start,start+ops):
                    col=get_column_letter(c);b,a,e=f'{col}{br+r}',f'{col}{ar+r}',f'{col}{er+r}'
                    v[b]=2;v[a]=1;v[e]=1;f[e]=f'={b}-{a}'
        v[f'{get_column_letter(sc)}27']=1.5
        return v,f
    def test_full_25_and_50_each_side(self):
        for line,(_,ops,_) in LAYOUT.items():
            v,f=self.fixture(line);r=parse_line(v,f,line,'hash')
            self.assertEqual(ops*10,len(r['points']));self.assertEqual('2026-09-01',r['reportDate'])
    def test_negative_quarantines_report(self):
        v,f=self.fixture('PT2');v['H14']=3
        with self.assertRaisesRegex(ValueError,'negative'):parse_line(v,f,'PT2','hash')
    def test_blank_is_not_zero(self):
        v,f=self.fixture('PT1');v['B4']=None
        with self.assertRaisesRegex(ValueError,'missing'):parse_line(v,f,'PT1','hash')
    def test_stale_cache_rejected(self):
        v,f=self.fixture('PT1');v['B33']=2
        with self.assertRaisesRegex(ValueError,'stale'):parse_line(v,f,'PT1','hash')
    def test_month_date_and_quarantine(self):
        v,f={},{}
        for line in LAYOUT:
            a,b=self.fixture(line);v.update(a);f.update(b)
        v['B14']=3
        reports,excluded=parse_sheet(v,f,'0902','hash')
        self.assertEqual(3,len(reports));self.assertEqual('PT1',excluded[0]['lineCode'])
        self.assertTrue(all(r['reportDate']=='2026-09-02' and r['provenance']['sheet']=='0902' for r in reports))
    def test_blank_sheet_and_outside_month(self):
        self.assertEqual(([],[]),parse_sheet({}, {}, '0930','hash'))
        with self.assertRaises(ValueError):parse_sheet({}, {}, '1001','hash')
    def test_incomplete_report_quarantined(self):
        v,f=self.fixture('PT1');v['B14']=None
        reports,excluded=parse_sheet(v,f,'0918','hash')
        self.assertEqual([],reports);self.assertIn('missing',excluded[0]['reason'])
if __name__=='__main__':unittest.main()
